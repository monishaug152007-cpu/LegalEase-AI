import os, html, re
from datetime import date
from io import BytesIO
import requests, streamlit as st
from dotenv import load_dotenv
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from fpdf import FPDF

load_dotenv()
API_URL=os.getenv("API_URL","http://127.0.0.1:8000").rstrip("/")
st.set_page_config(page_title="LegalEase",page_icon="⚖️",layout="centered")
st.markdown("""<style>
.stApp{background:#0e1117;color:#f1f5f9}.block-container{max-width:850px;padding-top:2rem}
h1,h2,h3{color:#f8fafc}
div[data-testid="stTextArea"] textarea,div[data-testid="stTextInput"] input{background:#202431;color:#f8fafc}
.preview{background:#171b26;border:1px solid #3b4354;border-radius:12px;padding:22px;white-space:pre-wrap;line-height:1.65;color:#edf2f7;max-height:520px;overflow:auto}
</style>""",unsafe_allow_html=True)
st.markdown("<h1 style='text-align:center'>⚖ LegalEase</h1><p style='text-align:center;color:#aab2c0'>AI-Powered Legal Document Generator</p>",unsafe_allow_html=True)
st.caption("Enter details, generate a draft, edit it, and export it.")
with st.form("legal_form"):
    doc_type=st.selectbox("Document Type",["Employment Contract","Lease Agreement","Non-Disclosure Agreement (NDA)","Freelance Work Contract","Employment Offer Letter","General Agreement","Other"])
    if doc_type=="Other": doc_type=st.text_input("Specify document type")
    parties=st.text_area("Parties Involved",placeholder="e.g., Jane Doe (Service Provider), ABC Ltd (Client)",height=90)
    terms=st.text_area("Terms & Conditions",placeholder="Separate clauses with semicolons (;)",height=130)
    effective_date=st.date_input("Effective Date",value=date.today())
    submitted=st.form_submit_button("✨ Generate Document",type="primary",use_container_width=True)
if submitted:
    if not doc_type or not parties.strip() or not terms.strip():
        st.error("Please complete Document Type, Parties Involved, and Terms & Conditions.")
    else:
        with st.spinner("Connecting to LegalEase AI..."):
            try:
                r=requests.post(f"{API_URL}/generate",json={"document_type":doc_type,"parties":parties.strip(),"terms":terms.strip(),"dates":effective_date.isoformat()},timeout=100)
                if r.ok:
                    st.session_state["generated_document"]=r.json()["document"]; st.success("Document generated successfully.")
                else: st.error(r.json().get("detail","Backend request failed."))
            except requests.RequestException as exc: st.error(f"Backend is not running at {API_URL}. Start backend terminal first. ({exc})")

def clean(s):
    return s.replace("\u2018","'").replace("\u2019","'").replace("\u201c",'"').replace("\u201d",'"').replace("\u2014","-").replace("\u2013","-")

def format_docx(text, title, terms):
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7)
    doc.styles["Normal"].font.name="Times New Roman"; doc.styles["Normal"].font.size=Pt(11)
    p=doc.add_paragraph(); p.alignment=1
    r=p.add_run("⚖ LegalEase"); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=RGBColor(30,55,90)
    p=doc.add_paragraph(); p.alignment=1; r=p.add_run(title); r.bold=True; r.font.size=Pt(14)
    for line in clean(text).splitlines():
        line=line.strip()
        if not line: continue
        if len(line)<100 and (line.isupper() or re.match(r"^\d+[.)]\\s+",line) or line.endswith(":")): doc.add_heading(line.rstrip(":"),level=2)
        else: doc.add_paragraph(line)
    clauses=[x.strip() for x in terms.split(";") if x.strip()]
    if clauses:
        doc.add_heading("Terms Provided",level=2); table=doc.add_table(rows=1,cols=2); table.style="Table Grid"
        table.rows[0].cells[0].text="No."; table.rows[0].cells[1].text="Term"
        for i,clause in enumerate(clauses,1):
            cells=table.add_row().cells; cells[0].text=str(i); cells[1].text=clause
    sec.footer.paragraphs[0].alignment=1; sec.footer.paragraphs[0].add_run("LegalEase | Draft for review only")
    out=BytesIO(); doc.save(out); return out.getvalue()

def format_pdf(text,title):
    pdf=FPDF(); pdf.set_auto_page_break(auto=True,margin=18); pdf.set_margins(18,18,18); pdf.alias_nb_pages(); pdf.add_page()
    pdf.set_font("Helvetica","B",16); pdf.cell(0,10,"LegalEase",new_x="LMARGIN",new_y="NEXT",align="C")
    pdf.set_font("Helvetica","B",13); pdf.multi_cell(0,8,clean(title),align="C"); pdf.ln(4); pdf.set_font("Helvetica",size=10)
    for line in clean(text).splitlines():
        line=line.strip()
        if not line: pdf.ln(2); continue
        if len(line)<100 and (line.isupper() or re.match(r"^\d+[.)]\\s+",line) or line.endswith(":")):
            pdf.ln(2); pdf.set_font("Helvetica","B",11); pdf.multi_cell(0,7,line); pdf.set_font("Helvetica",size=10)
        else: pdf.multi_cell(0,6,line)
    for page in range(1,pdf.pages_count+1):
        pdf.page=page; pdf.set_y(-14); pdf.set_font("Helvetica",size=8); pdf.cell(0,8,f"LegalEase | Draft for review | Page {page}/{{nb}}",align="C")
    return bytes(pdf.output())

if "generated_document" in st.session_state:
    st.subheader("Document Generated")
    edited=st.text_area("Edit Document Below",value=st.session_state["generated_document"],height=360,key="edited_document")
    st.markdown("**Preview**"); st.markdown(f'<div class="preview">{html.escape(edited)}</div>',unsafe_allow_html=True)
    filename=re.sub(r"[^a-zA-Z0-9]+","_",doc_type).strip("_").lower() or "legal_document"
    a,b,c=st.columns(3)
    with a: st.download_button("⬇ Download TXT",edited,file_name=f"{filename}.txt",mime="text/plain",use_container_width=True)
    with b: st.download_button("⬇ Download DOCX",format_docx(edited,doc_type,terms),file_name=f"{filename}.docx",mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",use_container_width=True)
    with c: st.download_button("⬇ Download PDF",format_pdf(edited,doc_type),file_name=f"{filename}.pdf",mime="application/pdf",use_container_width=True)
st.divider(); st.caption("AI-generated documents are drafts, not legal advice. Have a qualified lawyer review before signing. Avoid entering highly sensitive information.")
